#!/usr/bin/env python3
"""Build a unified AIC eval suite from the round Excel files in this folder.

Each case has:
  - query_id / query (content)
  - type: KIS | TRAKE | QA  (from the sheet name)
  - gt: parsed ranked list from column "Nộp"
        (falls back to "Agent" only when Nộp is empty)
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: uv pip install openpyxl") from exc

ROOT = Path(__file__).resolve().parent

SOURCE_FILES = [
    ("vong1", "Vòng 1.xlsx"),
    ("vong2", "Vòng 2.xlsx"),
    ("day3", "Day3.xlsx"),
]

VIDEO_RE = re.compile(r"^(L\d+_V\d+)\s*,\s*(.*)$", re.IGNORECASE)
VIDEO_ONLY_RE = re.compile(r"^(L\d+_V\d+)\s*$", re.IGNORECASE)
ANSWER_RE = re.compile(r""",\s*["'](.*)["']\s*$""")
EVENT_RE = re.compile(r"^E(\d+)\s*[:：.\-–]?\s*(.+)$", re.MULTILINE)
HEADER_ALIASES = {
    "query name": "query_id",
    "query": "query_id",
    "query_id": "query_id",
    "query id": "query_id",
    "description": "query",
    "nộp": "gt",
    "nop": "gt",
    "agent": "agent",
    "tham khảo": "reference",
    "tham khao": "reference",
    "note": "note",
    "link youtube": "youtube",
}


def _norm_header(value: Any) -> str:
    return str(value or "").strip().lower().replace("\xa0", " ")


def detect_header(ws) -> tuple[int, dict[str, int]]:
    for row in range(1, min(6, (ws.max_row or 1) + 1)):
        mapping: dict[str, int] = {}
        for col in range(1, (ws.max_column or 1) + 1):
            key = HEADER_ALIASES.get(_norm_header(ws.cell(row, col).value))
            if key:
                mapping[key] = col
        if "query_id" in mapping and "query" in mapping:
            return row, mapping
    raise ValueError(f"Cannot find header row in sheet {ws.title!r}")


def cell_text(ws, row: int, col: Optional[int]) -> str:
    if not col:
        return ""
    value = ws.cell(row, col).value
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def parse_gt_line(line: str) -> Optional[dict[str, Any]]:
    original = line.strip()
    line = original.strip(",")
    # Drop a single unmatched wrapping quote (Excel artifacts), keep QA "answer" quotes.
    if line.startswith('"') and line.count('"') == 1:
        line = line[1:]
    if line.endswith('"') and line.count('"') == 1:
        line = line[:-1]
    if line.startswith("'") and line.count("'") == 1:
        line = line[1:]
    if line.endswith("'") and line.count("'") == 1:
        line = line[:-1]
    line = line.strip()
    if not line:
        return None

    video_only = VIDEO_ONLY_RE.match(line)
    if video_only:
        return {
            "video_id": video_only.group(1),
            "frames": [],
            "frame_idx": None,
            "answer": None,
            "raw": line,
        }

    match = VIDEO_RE.match(line)
    if not match:
        return {"video_id": None, "frames": [], "frame_idx": None, "answer": None, "raw": line}

    video_id = match.group(1)
    rest = match.group(2).strip()
    answer = None
    answer_match = ANSWER_RE.search(rest)
    if answer_match:
        answer = answer_match.group(1)
        rest = rest[: answer_match.start()].rstrip()

    frames: list[int] = []
    for token in rest.split(","):
        token = token.strip().strip('"').strip("'")
        if re.fullmatch(r"-?\d+", token):
            frames.append(int(token))

    return {
        "video_id": video_id,
        "frames": frames,
        "frame_idx": frames[0] if frames else None,
        "answer": answer,
        "raw": line,
    }


def parse_gt_block(text: str) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for line in (text or "").splitlines():
        parsed = parse_gt_line(line)
        if parsed:
            items.append(parsed)
    return items


def parse_events(query: str) -> list[dict[str, str]]:
    events = []
    for match in EVENT_RE.finditer(query or ""):
        events.append({"id": f"E{match.group(1)}", "text": match.group(2).strip()})
    return events


def unique_keep_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            out.append(value)
    return out


def compact_gt(item: dict[str, Any]) -> str:
    video = item.get("video_id") or "?"
    frames = ",".join(str(f) for f in item.get("frames") or [])
    answer = item.get("answer")
    if answer is not None and frames:
        return f'{video}:{frames}:"{answer}"'
    if answer is not None:
        return f'{video}:"{answer}"'
    if frames:
        return f"{video}:{frames}"
    return video


def iter_cases(path: Path, source: str) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    cases: list[dict[str, Any]] = []
    for sheet_name in wb.sheetnames:
        qtype = sheet_name.strip().upper()
        if qtype not in {"KIS", "TRAKE", "QA"}:
            continue
        ws = wb[sheet_name]
        header_row, cols = detect_header(ws)
        for row in range(header_row + 1, (ws.max_row or 1) + 1):
            query_id = cell_text(ws, row, cols.get("query_id"))
            query = cell_text(ws, row, cols.get("query"))
            if not query_id and not query:
                continue
            nop = cell_text(ws, row, cols.get("gt"))
            agent = cell_text(ws, row, cols.get("agent"))
            if nop:
                gt_source, gt_raw = "Nộp", nop
            elif agent:
                gt_source, gt_raw = "Agent", agent
            else:
                gt_source, gt_raw = None, ""
            gt = parse_gt_block(gt_raw)
            case_id = f"{source}/{query_id or f'row-{row}'}"
            video_ids = unique_keep_order([x["video_id"] for x in gt if x.get("video_id")])
            answers = unique_keep_order([x["answer"] for x in gt if x.get("answer") is not None])
            cases.append(
                {
                    "id": case_id,
                    "query_id": query_id,
                    "query": query,
                    "type": qtype,
                    "source": source,
                    "source_file": path.name,
                    "sheet": sheet_name,
                    "gt_source": gt_source,
                    "gt": gt,
                    "gt_raw": gt_raw,
                    "n_gt": len(gt),
                    "video_ids": video_ids,
                    "answers": answers,
                    "events": parse_events(query),
                    "note": cell_text(ws, row, cols.get("note")),
                    "has_gt": bool(gt),
                }
            )
    return cases


def build_suite(folder: Path) -> dict[str, Any]:
    queries: list[dict[str, Any]] = []
    missing_files = []
    for source, filename in SOURCE_FILES:
        path = folder / filename
        if not path.is_file():
            missing_files.append(filename)
            continue
        queries.extend(iter_cases(path, source))

    if missing_files:
        raise FileNotFoundError(f"Missing Excel files: {missing_files}")

    by_source = Counter(q["source"] for q in queries)
    by_type = Counter(q["type"] for q in queries)
    missing_gt = [q["id"] for q in queries if not q["has_gt"]]
    fallback_agent = [q["id"] for q in queries if q["gt_source"] == "Agent"]

    return {
        "name": "AIC2025 Mentos evaluation suite",
        "version": "1.0",
        "created_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "sources": [{"id": sid, "file": fname} for sid, fname in SOURCE_FILES],
        "fields": {
            "query_id": "Excel column Query Name",
            "query": "Excel column Description",
            "gt": 'Excel column Nộp (fallback: Agent if Nộp is empty)',
            "type": "Sheet name: KIS | TRAKE | QA",
        },
        "counts": {
            "total": len(queries),
            "by_source": dict(by_source),
            "by_type": dict(by_type),
            "with_gt": sum(1 for q in queries if q["has_gt"]),
            "missing_gt": len(missing_gt),
            "gt_from_agent_fallback": len(fallback_agent),
        },
        "missing_gt_ids": missing_gt,
        "gt_from_agent_ids": fallback_agent,
        "queries": queries,
    }


def write_csv(path: Path, queries: list[dict[str, Any]]) -> None:
    fields = [
        "id",
        "query_id",
        "type",
        "source",
        "query",
        "gt",
        "n_gt",
        "video_ids",
        "answers",
        "gt_source",
        "has_gt",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for q in queries:
            writer.writerow(
                {
                    "id": q["id"],
                    "query_id": q["query_id"],
                    "type": q["type"],
                    "source": q["source"],
                    "query": q["query"],
                    "gt": " | ".join(compact_gt(item) for item in q["gt"]),
                    "n_gt": q["n_gt"],
                    "video_ids": " ".join(q["video_ids"]),
                    "answers": " | ".join(q["answers"]),
                    "gt_source": q["gt_source"] or "",
                    "has_gt": q["has_gt"],
                }
            )


def write_jsonl(path: Path, queries: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as fh:
        for q in queries:
            fh.write(json.dumps(q, ensure_ascii=False) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=ROOT)
    parser.add_argument("--out-dir", type=Path, default=ROOT)
    args = parser.parse_args()

    suite = build_suite(args.input_dir)
    out = args.out_dir
    out.mkdir(parents=True, exist_ok=True)

    json_path = out / "test_suite.json"
    jsonl_path = out / "test_suite.jsonl"
    csv_path = out / "test_suite.csv"

    json_path.write_text(json.dumps(suite, ensure_ascii=False, indent=2), encoding="utf-8")
    write_jsonl(jsonl_path, suite["queries"])
    write_csv(csv_path, suite["queries"])

    print(f"Wrote {json_path}")
    print(f"Wrote {jsonl_path}")
    print(f"Wrote {csv_path}")
    print("counts:", json.dumps(suite["counts"], ensure_ascii=False))
    if suite["missing_gt_ids"]:
        print("missing GT:", ", ".join(suite["missing_gt_ids"]))
    if suite["gt_from_agent_ids"]:
        print("GT fallback from Agent:", ", ".join(suite["gt_from_agent_ids"]))


if __name__ == "__main__":
    main()
